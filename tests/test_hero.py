from __future__ import annotations

import io
from datetime import date, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.importers import parse_upload
from app.main import ConfirmInput, app, normalize_import_snapshot
from app.scheduling import simulate, validate_tasks
from app.storage import Store
from app.worker import run_once


HERO = Path(__file__).resolve().parents[1] / "data" / "l1_project" / "hero_battery_factory_project.xlsx"


def hero_snapshot():
    parsed = parse_upload(HERO.name, HERO.read_bytes())
    snapshot = normalize_import_snapshot(parsed, {"name": "새 프로젝트", "mode": "REPLAY"}, ConfirmInput())
    return parsed, snapshot


def test_hero_import_preserves_schema_and_duration_semantics():
    parsed, snapshot = hero_snapshot()
    tasks = snapshot["tasks"]

    assert len(tasks) == 64
    assert [task["task_id"] for task in tasks] == [f"T{number:03d}" for number in range(1, 65)]
    assert all(task.get("name") for task in tasks)
    assert all(task.get("baseline_start") and task.get("baseline_finish") for task in tasks)
    assert all(task.get("owner") and task.get("location") for task in tasks)
    assert sum(task["dependency_type"] == "SS" for task in tasks) == 5
    assert {task["dependency_type"] for task in tasks} == {"FS", "SS"}
    assert validate_tasks(tasks) == []

    for source, task in zip(parsed["tasks"], tasks):
        start = date.fromisoformat(task["baseline_start"])
        finish = date.fromisoformat(task["baseline_finish"])
        assert (finish - start).days == source["duration_days"] == task["duration_days"]
        expected_workdays = sum(
            (start + timedelta(days=offset)).weekday() < 5
            for offset in range(task["duration_days"])
        )
        assert task["duration_workdays"] == expected_workdays
        assert task["duration_semantics"] == "calendar_days_elapsed"
        assert task["finish_boundary"] == "exclusive"


def test_hero_baseline_simulation_is_date_stable():
    _, snapshot = hero_snapshot()
    result = simulate(snapshot["project"], snapshot["tasks"])
    by_id = {task["task_id"]: task for task in result["schedule"]}

    assert result["violations"] == []
    assert result["finish_date"] == "2027-12-21"
    assert all(
        (task["baseline_start"], task["baseline_finish"])
        == (by_id[task["task_id"]]["planned_start"], by_id[task["task_id"]]["planned_finish"])
        for task in snapshot["tasks"]
    )


def test_ss_uses_predecessor_start_and_does_not_fall_back_to_fs():
    tasks = [
        {
            "task_id": "A", "name": "predecessor", "baseline_start": "2026-01-05",
            "baseline_finish": "2026-01-16", "duration_workdays": 10,
            "predecessor_ids": [], "dependency_type": "FS",
            "resource_demand": 1, "resource_capacity": 1,
        },
        {
            "task_id": "B", "name": "ss successor", "baseline_start": "2026-01-05",
            "baseline_finish": "2026-01-06", "duration_workdays": 2,
            "predecessor_ids": ["A"], "dependency_type": "SS",
            "resource_demand": 1, "resource_capacity": 1,
        },
    ]
    result = simulate({}, tasks)
    successor = next(task for task in result["schedule"] if task["task_id"] == "B")
    assert successor["planned_start"] == "2026-01-05"
    assert successor["planned_start"] < "2026-01-16"


@pytest.fixture
def hero_client(tmp_path, monkeypatch):
    monkeypatch.setenv("REPLAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REPLAN_DEMO_TOKEN", "test-token")
    monkeypatch.delenv("API_KEY", raising=False)
    return TestClient(app)


def _request(client, method, path, **kwargs):
    response = getattr(client, method)(path, headers={"Authorization": "Bearer test-token"}, **kwargs)
    assert response.status_code < 400, response.text
    return response.json()


def test_hero_api_smoke_upload_event_approval_commit_export(hero_client):
    project_id = _request(hero_client, "post", "/api/projects", json={"mode": "REPLAY"})["project_id"]
    preview = _request(
        hero_client, "post", f"/api/projects/{project_id}/imports",
        files={"file": (HERO.name, HERO.read_bytes())},
    )
    assert len(preview["tasks"]) == 64
    baseline = _request(hero_client, "post", f"/api/projects/{project_id}/imports/{preview['import_id']}/confirm", json={})

    event = _request(
        hero_client, "post", f"/api/projects/{project_id}/events",
        json={
            "content": "SMOKE TEST: T046 변경 예정일 2026-10-15",
            "mode": "REPLAY", "simulation_as_of": "2026-09-24T00:00:00+09:00",
            "patch": {"estimated_finish": {"T046": "2026-10-15"}},
        },
    )
    queued = _request(hero_client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event["event_id"]})
    assert run_once(Store())
    run = _request(hero_client, "get", f"/api/runs/{queued['run_id']}")
    assert run["run"]["status"] == "succeeded"
    scenario = next(item for item in run["scenarios"] if item["data"]["option_ids"] == [])
    prepared = _request(hero_client, "post", f"/api/scenarios/{scenario['id']}/prepare")
    for action in prepared["actions"]:
        _request(hero_client, "patch", f"/api/actions/{action['id']}", json={"state": "ACCEPTED"})
    _request(
        hero_client, "post", f"/api/scenarios/{scenario['id']}/approve",
        json={"actor": "프로젝트 운영팀", "confirmed_conditions": scenario["data"]["required_confirmations"]},
    )
    committed = _request(hero_client, "post", f"/api/scenarios/{scenario['id']}/commit")
    assert committed["version_id"] != baseline["version_id"]
    exported = hero_client.get(
        f"/api/projects/{project_id}/export?version_id={committed['version_id']}",
        headers={"Authorization": "Bearer test-token"},
    )
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    task_ids = {row[0] for row in workbook.active.iter_rows(min_row=3, values_only=True)}
    assert {"T001", "T046", "T064"} <= task_ids
