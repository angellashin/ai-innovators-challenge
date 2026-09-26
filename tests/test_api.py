from __future__ import annotations

import io
import hashlib
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.storage import Store
from app.worker import run_once


DEMO = Path(__file__).resolve().parents[1] / "REPLAN_demo_inputs.xlsx"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REPLAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REPLAN_DEMO_TOKEN", "test-token")
    monkeypatch.delenv("API_KEY", raising=False)
    return TestClient(app)


def request(client, method, path, **kwargs):
    response = getattr(client, method)(path, headers={"Authorization": "Bearer test-token"}, **kwargs)
    assert response.status_code < 400, response.text
    return response.json()


def baseline(client):
    project_id = request(client, "post", "/api/projects", json={"mode": "REPLAY"})["project_id"]
    preview = request(client, "post", f"/api/projects/{project_id}/imports", files={"file": ("demo.xlsx", DEMO.read_bytes())})
    assert len(preview["tasks"]) == 20
    confirmed = request(client, "post", f"/api/projects/{project_id}/imports/{preview['import_id']}/confirm", json={})
    return project_id, preview, confirmed


def test_original_upload_is_preserved(client):
    project_id = request(client, "post", "/api/projects", json={"mode": "REPLAY"})["project_id"]
    content = DEMO.read_bytes()
    preview = request(client, "post", f"/api/projects/{project_id}/imports", files={"file": ("demo.xlsx", content)})
    assert preview["_upload"]["sha256"] == hashlib.sha256(content).hexdigest()
    response = client.get(
        f"/api/projects/{project_id}/imports/{preview['import_id']}/original",
        headers={"Authorization": "Bearer test-token"},
    )
    assert response.status_code == 200 and response.content == content


def test_watch_plan_rejects_unapproved_source_host(client):
    project_id = request(client, "post", "/api/projects", json={"mode": "LIVE"})["project_id"]
    response = client.put(
        f"/api/projects/{project_id}/watch-plan",
        headers={"Authorization": "Bearer test-token"},
        json={"enabled": True, "source_allowlist": ["https://127.0.0.1/private"]},
    )
    assert response.status_code == 422


def test_analysis_requires_review_for_inferred_change(client):
    project_id, preview, _ = baseline(client)
    event = request(
        client, "post", f"/api/projects/{project_id}/events",
        json={"content": preview["events"][0]["body"], "source_label": "supplier@example.com"},
    )
    response = client.post(
        f"/api/projects/{project_id}/analyses",
        headers={"Authorization": "Bearer test-token"},
        json={"event_id": event["event_id"]},
    )
    assert response.status_code == 409
    request(client, "patch", f"/api/projects/{project_id}/events/{event['event_id']}/review", json={"confirmed": True})
    queued = request(client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event["event_id"]})
    assert queued["status"] == "queued"


def e01(client, project_id, preview):
    event = request(
        client, "post", f"/api/projects/{project_id}/events",
        json={
            "event_id": "E01", "content": preview["events"][0]["body"],
            "mode": "REPLAY", "data_origin": "SYNTHETIC",
            "simulation_as_of": "2026-09-23T09:00:00+02:00",
        },
    )
    request(client, "patch", f"/api/projects/{project_id}/events/{event['event_id']}/review", json={"confirmed": True})
    return event["event_id"]


def test_e01_budget_replan_approval_commit_and_export(client):
    project_id, preview, baseline_version = baseline(client)
    event_id = e01(client, project_id, preview)
    queued = request(client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event_id})
    assert run_once(Store())
    result = request(client, "get", f"/api/runs/{queued['run_id']}")
    assert result["run"]["status"] == "succeeded"
    assert result["run"]["data"]["agent_status"] == "llm_unavailable"
    scenarios = {tuple(item["data"]["option_ids"]): item for item in result["scenarios"]}
    assert scenarios[()]["data"]["budget_status"] == "UNSET"
    assert scenarios[()]["data"]["finish_date"] == "2026-10-30"
    assert scenarios[("OPT-02",)]["data"]["finish_date"] == "2026-10-29"
    assert scenarios[("OPT-03",)]["data"]["finish_date"] == "2026-10-28"

    bounded_run = request(client, "post", f"/api/runs/{queued['run_id']}/replan", json={"budget_krw": 3_000_000})
    assert run_once(Store())
    bounded = request(client, "get", f"/api/runs/{bounded_run['run_id']}")
    bounded_scenarios = {tuple(item["data"]["option_ids"]): item for item in bounded["scenarios"]}
    assert not bounded_scenarios[("OPT-03",)]["data"]["budget_met"]
    assert not any(item["data"]["target_met"] and item["data"]["budget_met"] for item in bounded["scenarios"])

    new_run = request(client, "post", f"/api/runs/{queued['run_id']}/replan", json={"budget_krw": 6000000})
    assert run_once(Store())
    updated = request(client, "get", f"/api/runs/{new_run['run_id']}")
    option = next(item for item in updated["scenarios"] if item["data"]["option_ids"] == ["OPT-03"])
    assert option["data"]["budget_met"] and option["data"]["target_met"]
    scenario_id = option["id"]

    unapproved = client.post(f"/api/scenarios/{scenario_id}/commit", headers={"Authorization": "Bearer test-token"})
    assert unapproved.status_code == 409
    prepared = request(client, "post", f"/api/scenarios/{scenario_id}/prepare")
    assert len(prepared["actions"]) == 1
    assert prepared["actions"][0]["data"]["due_at"] == "2026-09-11"
    assert prepared["actions"][0]["data"]["owner"] == "프로젝트 운영팀"
    pending = client.post(
        f"/api/scenarios/{scenario_id}/approve", headers={"Authorization": "Bearer test-token"},
        json={"actor": "프로젝트 운영팀", "confirmed_conditions": option["data"]["required_confirmations"]},
    )
    assert pending.status_code == 409
    request(client, "patch", f"/api/actions/{prepared['actions'][0]['id']}", json={"state": "ACCEPTED"})
    request(client, "post", f"/api/scenarios/{scenario_id}/approve", json={"actor": "프로젝트 운영팀", "confirmed_conditions": option["data"]["required_confirmations"]})
    request(client, "patch", f"/api/actions/{prepared['actions'][0]['id']}", json={"state": "REJECTED"})
    revoked = client.post(f"/api/scenarios/{scenario_id}/commit", headers={"Authorization": "Bearer test-token"})
    assert revoked.status_code == 409
    request(client, "patch", f"/api/actions/{prepared['actions'][0]['id']}", json={"state": "ACCEPTED"})
    committed = request(client, "post", f"/api/scenarios/{scenario_id}/commit")
    assert committed["version_id"] != baseline_version["version_id"]
    assert request(client, "post", f"/api/scenarios/{scenario_id}/commit")["duplicate"]

    state = request(client, "get", f"/api/projects/{project_id}")
    assert [item["status"] for item in state["versions"]] == ["committed", "baseline"]
    assert state["versions"][0]["scenario_id"] == scenario_id
    assert state["versions"][0]["parent_id"] == baseline_version["version_id"]
    assert [(item["scenario_id"], item["event_id"], item["run_id"]) for item in state["approvals"]] == [
        (scenario_id, event_id, new_run["run_id"])]

    response = client.get(f"/api/projects/{project_id}/export?version_id={committed['version_id']}", headers={"Authorization": "Bearer test-token"})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert book.active["D1"].value == committed["version_hash"]
    task_rows = {row[0]: row for row in book.active.iter_rows(min_row=3, values_only=True)}
    assert task_rows["T20"][5] == "2026-10-28"


def test_analysis_uses_queued_project_context_snapshot(client):
    project_id, preview, _ = baseline(client)
    event_id = e01(client, project_id, preview)
    request(
        client,
        "post",
        f"/api/projects/{project_id}/supplier-calendars",
        json={"supplier_id": "SUP-01", "label": "공급사 휴무", "unavailable_dates": ["2026-10-05"]},
    )
    queued = request(client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event_id})
    before = request(client, "get", f"/api/runs/{queued['run_id']}")
    snapshot = before["run"]["data"]["project_context_snapshot"]
    assert "2026-10-05" in snapshot["project"]["supplier_unavailable_dates"]
    request(
        client,
        "post",
        f"/api/projects/{project_id}/supplier-calendars",
        json={"supplier_id": "SUP-02", "label": "변경 후 휴무", "unavailable_dates": ["2026-10-06"]},
    )
    assert run_once(Store())
    after = request(client, "get", f"/api/runs/{queued['run_id']}")
    run_snapshot = after["run"]["data"]["project_context_snapshot"]
    assert "2026-10-06" not in run_snapshot["project"]["supplier_unavailable_dates"]
    assert {item["data"]["project_context_hash"] for item in after["scenarios"]} == {snapshot["content_hash"]}


def test_duplicate_event_and_stale_scenario(client):
    project_id, preview, _ = baseline(client)
    event_id = e01(client, project_id, preview)
    repeated = request(client, "post", f"/api/projects/{project_id}/events", json={"event_id": "E01", "content": preview["events"][0]["body"], "mode": "REPLAY", "data_origin": "SYNTHETIC", "simulation_as_of": "2026-09-23T09:00:00+02:00"})
    assert repeated["duplicate"] and repeated["event_id"] == event_id
    queued = request(client, "post", f"/api/projects/{project_id}/analyses", json={"event_id": event_id, "budget_krw": 6000000})
    assert run_once(Store())
    scenarios = request(client, "get", f"/api/runs/{queued['run_id']}")["scenarios"]
    first = next(item for item in scenarios if item["data"]["option_ids"] == ["OPT-03"])
    second = next(item for item in scenarios if item["data"]["option_ids"] == ["OPT-02"])
    for item in (first, second):
        actions = request(client, "post", f"/api/scenarios/{item['id']}/prepare")["actions"]
        for action in actions:
            request(client, "patch", f"/api/actions/{action['id']}", json={"state": "ACCEPTED"})
        request(client, "post", f"/api/scenarios/{item['id']}/approve", json={"actor": "프로젝트 운영팀", "confirmed_conditions": item["data"]["required_confirmations"]})
    request(client, "post", f"/api/scenarios/{first['id']}/commit")
    conflict = client.post(f"/api/scenarios/{second['id']}/commit", headers={"Authorization": "Bearer test-token"})
    assert conflict.status_code == 409


def test_revised_excel_requires_diff_confirmation(client):
    project_id, _, _ = baseline(client)
    workbook = load_workbook(DEMO)
    workbook["일정표"]["E10"] = "2026-09-30"
    changed = io.BytesIO()
    workbook.save(changed)
    preview = request(client, "post", f"/api/projects/{project_id}/imports", files={"file": ("revised.xlsx", changed.getvalue())})
    assert preview["import_kind"] == "change"
    assert preview["diff"]["summary"]["changed"] >= 1
    confirmed = request(client, "post", f"/api/projects/{project_id}/imports/{preview['import_id']}/confirm", json={})
    assert confirmed["event"]["channel"] == "revised_excel"
    assert "T03" in confirmed["event"]["related_task_ids"]


def test_demo_auth_required(client):
    assert client.get("/api/projects").status_code == 401


def test_p1_document_mail_notifications_and_site_prep(client):
    project_id, preview, _ = baseline(client)
    mail = request(client, "put", f"/api/projects/{project_id}/mail-account", json={"provider": "imap", "host": "imap.example.com", "username": "ops@example.com"})
    assert mail["account"]["credentials_required"] is True
    uploaded = request(client, "post", f"/api/projects/{project_id}/documents", files={"file": ("notice.txt", "T03 공급사 일정이 하루 지연됩니다.".encode())})
    assert uploaded["status"] == "queued"
    assert uploaded["document"]["status"] == "QUEUED"
    assert run_once(Store())
    processed = request(client, "get", f"/api/projects/{project_id}/documents/{uploaded['document_id']}")
    assert processed["document"]["data"]["status"] == "SUCCEEDED"
    assert processed["document"]["data"]["input_type"] == "text"
    assert processed["run"]["status"] == "succeeded"
    assert processed["run"]["data"]["event_id"]
    duplicate = request(client, "post", f"/api/projects/{project_id}/documents", files={"file": ("notice-copy.txt", "T03 공급사 일정이 하루 지연됩니다.".encode())})
    assert duplicate["duplicate"] is True
    assert duplicate["document_id"] == uploaded["document_id"]
    assert len(request(client, "get", f"/api/projects/{project_id}/documents")["documents"]) == 1
    notifications = request(client, "get", f"/api/projects/{project_id}/notifications")
    assert notifications["notifications"]
    channel = request(client, "post", f"/api/projects/{project_id}/notification-channels", json={"channel": "email", "target": "ops@example.com"})
    assert channel["channel"]["delivery_mode"] == "DRAFT"
    prep = request(client, "post", f"/api/projects/{project_id}/site-prep", json={})
    assert len(prep["items"]) == 4
    assert request(client, "get", f"/api/projects/{project_id}/documents")["documents"]


def test_p1_supplier_calendar_and_feed_registration(client, monkeypatch):
    project_id = request(client, "post", "/api/projects", json={"mode": "LIVE"})["project_id"]
    monkeypatch.setenv("REPLAN_ALLOWED_SOURCE_HOSTS", "environment.ec.europa.eu")
    feed = request(client, "post", f"/api/projects/{project_id}/public-feeds", json={"label": "EU news", "url": "https://environment.ec.europa.eu/news_en"})
    assert feed["feed"]["kind"] == "rss"
    calendar = request(client, "post", f"/api/projects/{project_id}/supplier-calendars", json={"supplier_id": "SUP-01", "label": "공급사 휴무", "unavailable_dates": ["2026-10-03"]})
    assert "2026-10-03" in calendar["supplier_unavailable_dates"]
    project = request(client, "get", f"/api/projects/{project_id}")
    assert "2026-10-03" in project["project"]["supplier_unavailable_dates"]


def test_p1_rejects_malformed_pdf(client):
    project_id, _, _ = baseline(client)
    response = client.post(
        f"/api/projects/{project_id}/documents",
        headers={"Authorization": "Bearer test-token"},
        files={"file": ("malformed.pdf", b"%PDF-1.7\nnot-a-valid-pdf")},
    )
    assert response.status_code == 202
    uploaded = response.json()
    assert uploaded["document"]["status"] == "QUEUED"
    assert run_once(Store())
    processed = client.get(
        f"/api/projects/{project_id}/documents/{uploaded['document_id']}",
        headers={"Authorization": "Bearer test-token"},
    )
    assert processed.status_code == 200
    assert processed.json()["document"]["data"]["status"] == "FAILED"
    assert processed.json()["run"]["status"] == "failed"
    retried = request(client, "post", f"/api/projects/{project_id}/documents/{uploaded['document_id']}/retry")
    assert retried["status"] == "queued"
    assert retried["document"]["retry_count"] == 1
    assert run_once(Store())
    failed_again = request(client, "get", f"/api/projects/{project_id}/documents/{uploaded['document_id']}")
    assert failed_again["document"]["data"]["status"] == "FAILED"
    assert failed_again["document"]["data"]["retry_count"] == 1
