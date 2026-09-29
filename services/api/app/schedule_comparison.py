"""Immutable schedule comparisons shared by the UI and the exported workbook."""
from datetime import date, timedelta
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def comparison(db, project_id, version, scenario=None):
    baseline = version
    while baseline.get('parent_id'):
        baseline = db.get_json('versions', baseline['parent_id'], project_id)
    if scenario is None and version['data'].get('scenario_id'):
        scenario = db.get_json('scenarios', version['data']['scenario_id'], project_id)
    current = version['data']['tasks']
    risk = None
    if scenario:
        with db.connection() as conn:
            ids = conn.execute('SELECT id FROM scenarios WHERE project_id=? AND run_id=?', (project_id, scenario['run_id'])).fetchall()
        candidates = [db.get_json('scenarios', row['id'], project_id) for row in ids]
        risk = next((s for s in candidates if not s['data'].get('option_ids')), None)
        current = scenario['data']['schedule']
    def dates(task):
        return {'start': task.get('planned_start') or task.get('baseline_start'), 'finish': task.get('planned_finish') or task.get('baseline_finish')}
    def index(tasks):
        return {t['task_id']: dates(t) for t in tasks}
    old = index(baseline['data']['tasks'])
    new = index(current)
    exposed = index(risk['data']['schedule']) if risk else {}
    def delta(a, b):
        return (date.fromisoformat(a[:10]) - date.fromisoformat(b[:10])).days if a and b else None
    rows = []
    for task in baseline['data']['tasks']:
        key = task['task_id']
        if key not in new:
            continue
        a, b, c = old[key], exposed.get(key), new[key]
        rows.append({'task_id': key, 'name': task.get('name', key), 'baseline': a, 'risk': b, 'revised': c,
                     'shift_days': delta(c['finish'], a['finish']), 'changed': a != c or (b is not None and b != a)})
    def finish(values):
        return max((v['finish'] for v in values.values() if v.get('finish')), default=None)
    a, b, c = finish(old), finish(exposed), finish(new)
    return {'baseline_version_id': baseline['id'], 'version_id': version['id'],
            'scenario_id': scenario['id'] if scenario else None,
            'committed': bool(scenario and version['status'] == 'committed' and version['data'].get('scenario_id') == scenario['id']),
            'baseline_finish': a, 'risk_finish': b, 'revised_finish': c,
            'risk_days': delta(b, a), 'recovered_days': delta(b, c), 'remaining_days': delta(c, a),
            'rows': rows, 'changed_count': sum(row['changed'] for row in rows)}


def add_comparison_sheets(wb, data, safe):
    ws = wb.create_sheet('일정 변화 요약', 1)
    ws.append(['일정 비교', '완료일', '일수'])
    ws.append(['최초 기준', data['baseline_finish'], 0])
    ws.append(['대응하지 않으면', data['risk_finish'], data['risk_days']])
    ws.append(['승인 반영 일정' if data['committed'] else '현재 일정', data['revised_finish'], data['remaining_days']])
    ws.append(['대응으로 회복', '', data['recovered_days']])
    ws.append(['해석', '전체 완료일의 달력 일수 차이. 개별 작업 지연은 겹치므로 더하지 않습니다.'])
    ws.append(['원본 버전', data['baseline_version_id']])
    ws.append(['반영 버전', data['version_id']])
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 85
    ws.column_dimensions['C'].width = 16
    chart = wb.create_sheet('전후 비교 간트', 2)
    changed = [r for r in data['rows'] if r['changed']] or data['rows']
    dates = [date.fromisoformat(v[k][:10]) for r in changed for t in ('baseline','risk','revised') if (v := r[t]) for k in ('start','finish') if v.get(k)]
    if not dates:
        return
    start, end = min(dates), max(dates)
    # Weekly bins: precise dates are preserved in the first columns.
    ticks = [start + timedelta(days=i) for i in range(0, (end-start).days+1, 7)]
    chart.append(['변경 영향 작업 · 주 단위 표시 / 정확한 날짜는 시작·종료 열 참조'])
    chart.append(['작업', '작업명', '구분', '시작', '종료', '최초 대비 종료 차이'] + [t.isoformat() for t in ticks])
    colors = {'baseline':'94A3B8', 'risk':'D97762', 'revised':'168C79'}
    labels = {'baseline':'최초 기준', 'risk':'무대응', 'revised':'승인 반영' if data['committed'] else '현재 일정'}
    for row in changed:
        for kind in colors:
            value = row[kind]
            if not value:
                continue
            chart.append([safe(row['task_id']), safe(row['name']), labels[kind], value['start'], value['finish'], row['shift_days'] if kind == 'revised' else None])
            a,b = date.fromisoformat(value['start'][:10]), date.fromisoformat(value['finish'][:10])
            for col,tick in enumerate(ticks,7):
                if a <= tick + timedelta(days=6) and b >= tick:
                    chart.cell(chart.max_row,col).fill = PatternFill('solid', fgColor=colors[kind])
    chart.freeze_panes = 'G3'
    for col,width in [('A',12),('B',46),('C',16),('D',14),('E',14),('F',20)]:
        chart.column_dimensions[col].width = width
    for col in range(7,7+len(ticks)):
        chart.column_dimensions[get_column_letter(col)].width = 12
    for sheet in [ws, chart]:
        for cell in sheet[1 if sheet == ws else 2]:
            cell.fill = PatternFill('solid', fgColor='172033')
            cell.font = Font(color='FFFFFF', bold=True)
        sheet.sheet_view.showGridLines = False
