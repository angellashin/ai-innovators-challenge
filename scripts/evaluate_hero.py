"""Leakage-safe, no-network Hero integration and benchmark evaluator.

This harness is intentionally independent from evaluate_replay.py. It evaluates
only outputs the current deterministic RE:PLAN pipeline actually produces and
marks unavailable capabilities as NOT_IMPLEMENTED.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "services" / "api"))

from app.events import normalize_event  # noqa: E402
from app.importers import parse_upload  # noqa: E402
from app.main import ConfirmInput, app, normalize_import_snapshot  # noqa: E402
from app.scheduling import simulate, validate_tasks  # noqa: E402
from app.storage import Store  # noqa: E402
from app.worker import run_once  # noqa: E402
from app.risk_signals import mentioned_risk_types, search_risk_signals  # noqa: E402
from app.task_retrieval import retrieve_related_tasks  # noqa: E402
from scripts.evaluation_mocks import MockEvaluationGateway  # noqa: E402


def _load(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _scores(predicted: set[str], expected: set[str]) -> dict[str, float | int]:
    true_positive = len(predicted & expected)
    precision = true_positive / len(predicted) if predicted else 0.0
    recall = true_positive / len(expected) if expected else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "true_positive": true_positive,
        "predicted_count": len(predicted),
        "expected_count": len(expected),
        "precision": round(precision, 6),
        "recall": round(recall, 6),
        "f1": round(f1, 6),
    }


def _aggregate(rows: list[dict[str, Any]]) -> dict[str, Any]:
    if not rows:
        return {"case_count": 0, "precision": 0.0, "recall": 0.0, "f1": 0.0}
    tp = sum(row["scores"]["true_positive"] for row in rows)
    predicted = sum(row["scores"]["predicted_count"] for row in rows)
    expected = sum(row["scores"]["expected_count"] for row in rows)
    precision = tp / predicted if predicted else 0.0
    recall = tp / expected if expected else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "case_count": len(rows),
        "micro": {"precision": round(precision, 6), "recall": round(recall, 6), "f1": round(f1, 6)},
        "macro": {
            key: round(sum(row["scores"][key] for row in rows) / len(rows), 6)
            for key in ("precision", "recall", "f1")
        },
    }


def _public_event_input(case: dict[str, Any]) -> str:
    """Build the agent input from non-answer fields only."""
    return "\n".join(
        value for value in (
            str(case.get("project_context") or "").strip(),
            str(((case.get("risk") or {}).get("cause") or {}).get("value") or "").strip(),
        ) if value
    )


def _leakage_audit(public_inputs: dict[str, str], mappings: list[dict[str, Any]]) -> dict[str, Any]:
    mapping_by_id = {item["case_id"]: item for item in mappings}
    violations = []
    for case_id, text in public_inputs.items():
        task_ids = set(re.findall(r"\bT\d{3}\b", text))
        hidden_ids = set(mapping_by_id[case_id].get("direct_task_ids") or []) | set(mapping_by_id[case_id].get("downstream_task_ids") or [])
        if task_ids & hidden_ids:
            violations.append({"case_id": case_id, "leaked_task_ids": sorted(task_ids & hidden_ids)})
    return {
        "pass": not violations,
        "violations": violations,
        "agent_input_fields": ["project_context", "risk.cause.value"],
        "withheld_fields": [
            "direct_task_ids", "downstream_task_ids", "risk.category", "schedule_impact",
            "response", "mapping.risk_category",
        ],
    }


def _run_api_smoke(project_path: Path) -> dict[str, Any]:
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    artifact_dir = PROJECT_ROOT / "artifacts"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    previous_dir = os.environ.get("REPLAN_DATA_DIR")
    previous_token = os.environ.get("REPLAN_DEMO_TOKEN")
    previous_key = os.environ.pop("API_KEY", None)
    try:
        with tempfile.TemporaryDirectory(prefix="hero-smoke-", dir=artifact_dir) as temp_dir:
            os.environ["REPLAN_DATA_DIR"] = temp_dir
            os.environ["REPLAN_DEMO_TOKEN"] = "hero-smoke-token"
            client = TestClient(app)
            headers = {"Authorization": "Bearer hero-smoke-token"}

            def request(method: str, path: str, **kwargs: Any) -> dict[str, Any]:
                response = getattr(client, method)(path, headers=headers, **kwargs)
                if response.status_code >= 400:
                    raise RuntimeError(f"{method.upper()} {path}: {response.status_code} {response.text}")
                return response.json()

            project_id = request("post", "/api/projects", json={"mode": "REPLAY"})["project_id"]
            preview = request(
                "post", f"/api/projects/{project_id}/imports",
                files={"file": (project_path.name, project_path.read_bytes())},
            )
            baseline = request("post", f"/api/projects/{project_id}/imports/{preview['import_id']}/confirm", json={})
            event = request(
                "post", f"/api/projects/{project_id}/events",
                json={
                    "content": "SMOKE TEST: T046 revised finish 2027-02-15",
                    "mode": "REPLAY", "simulation_as_of": "2026-09-24T00:00:00+09:00",
                    "patch": {"estimated_finish": {"T046": "2027-02-15"}},
                },
            )
            queued = request("post", f"/api/projects/{project_id}/analyses", json={"event_id": event["event_id"]})
            if not run_once(Store()):
                raise RuntimeError("analysis worker did not claim the smoke run")
            run = request("get", f"/api/runs/{queued['run_id']}")
            scenario = next(item for item in run["scenarios"] if item["data"]["option_ids"] == [])
            prepared = request("post", f"/api/scenarios/{scenario['id']}/prepare")
            for action in prepared["actions"]:
                request("patch", f"/api/actions/{action['id']}", json={"state": "ACCEPTED"})
            request(
                "post", f"/api/scenarios/{scenario['id']}/approve",
                json={"actor": "Hero smoke evaluator", "confirmed_conditions": scenario["data"]["required_confirmations"]},
            )
            committed = request("post", f"/api/scenarios/{scenario['id']}/commit")
            exported = client.get(
                f"/api/projects/{project_id}/export?version_id={committed['version_id']}", headers=headers,
            )
            if exported.status_code != 200:
                raise RuntimeError(f"export failed: {exported.status_code} {exported.text}")
            workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
            exported_ids = {row[0] for row in workbook.active.iter_rows(min_row=3, values_only=True) if row[0]}
            changed = [
                task["task_id"] for task in scenario["data"]["schedule"]
                if task.get("baseline_start") != task.get("planned_start")
                or task.get("baseline_finish") != task.get("planned_finish")
            ]
            return {
                "status": "PASS",
                "flow": ["xlsx_upload", "task_parsing", "event_application", "dependency_propagation", "approval", "revised_xlsx_generation"],
                "parsed_task_count": len(preview["tasks"]),
                "event_task_id": "T046",
                "changed_task_count": len(changed),
                "changed_task_ids": changed,
                "scenario_finish_date": scenario["data"]["finish_date"],
                "exported_task_count": len(exported_ids),
                "version_created": committed["version_id"] != baseline["version_id"],
            }
    finally:
        if previous_dir is None:
            os.environ.pop("REPLAN_DATA_DIR", None)
        else:
            os.environ["REPLAN_DATA_DIR"] = previous_dir
        if previous_token is None:
            os.environ.pop("REPLAN_DEMO_TOKEN", None)
        else:
            os.environ["REPLAN_DEMO_TOKEN"] = previous_token
        if previous_key is not None:
            os.environ["API_KEY"] = previous_key


def evaluate(
    project_path: Path,
    ground_truth_path: Path,
    mapping_path: Path,
    eval_dir: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    evaluation_files = sorted(eval_dir.glob("*.json"))
    evaluation_manifests = [
        {"file": path.name, "top_level_keys": sorted(_load(path).keys())}
        for path in evaluation_files
    ]
    parsed = parse_upload(project_path.name, project_path.read_bytes())
    snapshot = normalize_import_snapshot(parsed, {"name": "새 프로젝트", "mode": "REPLAY"}, ConfirmInput())
    project, tasks = snapshot["project"], snapshot["tasks"]
    validation_errors = validate_tasks(tasks)
    baseline = simulate(project, tasks)
    baseline_by_id = {task["task_id"]: task for task in baseline["schedule"]}
    changed_baseline = [
        task["task_id"] for task in tasks
        if (task["baseline_start"], task["baseline_finish"])
        != (baseline_by_id[task["task_id"]]["planned_start"], baseline_by_id[task["task_id"]]["planned_finish"])
    ]

    source_by_id = {task["task_id"]: task for task in parsed["tasks"]}
    duration_mismatches = []
    for task in tasks:
        start = date.fromisoformat(task["baseline_start"])
        finish = date.fromisoformat(task["baseline_finish"])
        if (finish - start).days != source_by_id[task["task_id"]].get("duration_days"):
            duration_mismatches.append(task["task_id"])

    integration = {
        "mode": "SMOKE_INTEGRATION",
        "project": str(project_path),
        "parsed_task_count": len(tasks),
        "task_ids_preserved": [task["task_id"] for task in tasks] == [f"T{number:03d}" for number in range(1, 65)],
        "field_coverage": {
            "name": sum(bool(task.get("name")) for task in tasks),
            "baseline_start": sum(bool(task.get("baseline_start")) for task in tasks),
            "baseline_finish": sum(bool(task.get("baseline_finish")) for task in tasks),
            "owner": sum(bool(task.get("owner")) for task in tasks),
            "location": sum(bool(task.get("location")) for task in tasks),
        },
        "dependency_types": {
            kind: sum(task.get("dependency_type") == kind for task in tasks) for kind in ("FS", "SS")
        },
        "ss_task_ids": [task["task_id"] for task in tasks if task.get("dependency_type") == "SS"],
        "predecessor_integrity": all(
            predecessor in {item["task_id"] for item in tasks}
            for task in tasks for predecessor in task.get("predecessor_ids") or []
        ),
        "duration_conversion": {
            "source_semantics": "calendar_days_elapsed; planned_end - planned_start == duration_days",
            "canonical_semantics": "duration_workdays counted on [baseline_start, baseline_finish)",
            "source_calendar_duration_mismatches": duration_mismatches,
            "converted_task_count": sum(task.get("duration_semantics") == "calendar_days_elapsed" for task in tasks),
            "finish_boundary_offsets_preserved": sum("finish_boundary_offset_days" in task for task in tasks),
        },
        "simulator_validation_errors": validation_errors,
        "baseline_simulation": {
            "finish_date": baseline["finish_date"],
            "violations": baseline["violations"],
            "changed_from_source_count": len(changed_baseline),
            "changed_task_ids": changed_baseline,
        },
        "api_smoke": _run_api_smoke(project_path),
    }

    ground_truth = _load(ground_truth_path)
    mapping_doc = _load(mapping_path)
    mappings = mapping_doc.get("mappings") or []
    eligible = [
        item for item in mappings
        if str(item.get("mapping_confidence") or "").lower() in {"high", "medium"}
        and item.get("direct_task_ids")
    ]
    excluded = [item for item in mappings if item not in eligible]
    cases_by_id = {case["case_id"]: case for case in ground_truth.get("cases") or []}
    public_inputs = {item["case_id"]: _public_event_input(cases_by_id[item["case_id"]]) for item in eligible}

    direct_rows = []
    agent_rows = []
    l2_exclusion_audit = []
    mock_gateway = MockEvaluationGateway()
    for mapping in eligible:
        case_id = mapping["case_id"]
        event = normalize_event({"content": public_inputs[case_id], "mode": "REPLAY"}, project, tasks)
        predicted = set(event.get("related_task_ids") or [])
        expected = set(mapping["direct_task_ids"])
        original_risk_id = str(cases_by_id[case_id].get("source_risk_id") or "")
        risk_types = mentioned_risk_types(public_inputs[case_id])
        analogous = search_risk_signals(risk_type=risk_types[0] if risk_types else "",
                                         exclude_source_risk_id=original_risk_id, limit=5)["results"]
        exposed = [item["risk_id"] for item in analogous]
        l2_exclusion_audit.append({"case_id": case_id, "excluded_source_risk_id": original_risk_id,
                                   "returned_risk_ids": exposed, "pass": original_risk_id not in exposed})
        agent_predicted = predicted or set(retrieve_related_tasks(public_inputs[case_id], tasks, mock_gateway, analogous)["task_ids"])
        direct_rows.append({
            "case_id": case_id,
            "classification_status": event.get("classification_status"),
            "predicted_task_ids": sorted(predicted),
            "expected_task_ids": sorted(expected),
            "scores": _scores(predicted, expected),
        })
        agent_rows.append({"case_id": case_id, "predicted_task_ids": sorted(agent_predicted),
                           "expected_task_ids": sorted(expected), "scores": _scores(agent_predicted, expected)})

    propagation_rows = []
    for mapping in eligible:
        direct_ids = set(mapping["direct_task_ids"])
        patch = {
            task_id: (date.fromisoformat(source_by_id[task_id]["planned_finish"]) + timedelta(days=7)).isoformat()
            for task_id in direct_ids
        }
        result = simulate(project, tasks, event={"event_id": f"PROP-{mapping['case_id']}", "patch": {"estimated_finish": patch}})
        result_by_id = {task["task_id"]: task for task in result["schedule"]}
        changed = {
            task_id for task_id, task in result_by_id.items()
            if task_id not in direct_ids
            and (task["planned_start"], task["planned_finish"])
            != (baseline_by_id[task_id]["planned_start"], baseline_by_id[task_id]["planned_finish"])
        }
        expected = set(mapping.get("downstream_task_ids") or [])
        propagation_rows.append({
            "case_id": mapping["case_id"],
            "seed": "hidden ground-truth direct_task_ids with controlled +7 calendar-day finish shift",
            "changed_downstream_task_ids": sorted(changed),
            "expected_downstream_task_ids": sorted(expected),
            "project_finish_shift_days": (
                date.fromisoformat(result["finish_date"]) - date.fromisoformat(baseline["finish_date"])
            ).days,
            "scores": _scores(changed, expected),
        })

    evaluation = {
        "mode": "BENCHMARK",
        "agent_pipeline": "app.events.normalize_event + quote-validated mocked task retrieval + app.scheduling.simulate",
        "paid_calls": 0,
        "eval_dir": str(eval_dir),
        "evaluation_files": evaluation_manifests,
        "eligibility": {
            "mapping_total": len(mappings),
            "main_benchmark_case_count": len(eligible),
            "included_case_ids": [item["case_id"] for item in eligible],
            "excluded_low_or_empty_count": len(excluded),
            "excluded_case_ids": [item["case_id"] for item in excluded],
            "rule": "mapping_confidence in {high, medium} and direct_task_ids non-empty",
        },
        "ground_truth_leakage_audit": _leakage_audit(public_inputs, mappings),
        "metrics": {
            "affected_task_retrieval": {
                "status": "IMPLEMENTED",
                "ranking_metrics": "NOT_AVAILABLE: service output has no explicit task ranking",
                "aggregate": _aggregate(direct_rows),
                "cases": direct_rows,
                "rules_only": _aggregate(direct_rows),
                "agent_included_mock": _aggregate(agent_rows),
                "agent_cases_mock": agent_rows,
                "mock_note": "Offline semantic fixture checks the agent interface, not actual LLM performance.",
                "l2_exclusion_audit": {"pass": all(row["pass"] for row in l2_exclusion_audit),
                                       "cases": l2_exclusion_audit},
            },
            "schedule_propagation": {
                "status": "IMPLEMENTED",
                "interpretation": "oracle-seeded dependency propagation; not delay prediction",
                "aggregate": _aggregate(propagation_rows),
                "cases": propagation_rows,
            },
            "delay_estimation": {
                "status": "NOT_IMPLEMENTED",
                "reason": "current event pipeline does not estimate delay magnitude from the leak-free L3 narrative",
            },
            "risk_classification": {
                "status": "NOT_IMPLEMENTED",
                "reason": "classification_status is schedule interpretation state, not a risk_category prediction",
            },
            "evidence_retrieval": {
                "status": "NOT_EVALUATED",
                "reason": "L2 search is available; this benchmark has no independently labeled relevance set for retrieval scoring",
            },
            "response_quality": {
                "status": "NOT_IMPLEMENTED",
                "reason": "deterministic pipeline returns no mitigation recommendation suitable for a rubric; no LLM judge configured",
            },
        },
    }
    return integration, evaluation


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--ground-truth", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--eval-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--integration-output", type=Path, default=PROJECT_ROOT / "artifacts" / "hero_integration_report.json")
    args = parser.parse_args()

    integration, evaluation = evaluate(
        args.project.resolve(), args.ground_truth.resolve(), args.mapping.resolve(), args.eval_dir.resolve()
    )
    for path, payload in ((args.integration_output, integration), (args.output, evaluation)):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(evaluation, ensure_ascii=False, indent=2))
    passed = (
        not integration["simulator_validation_errors"]
        and integration["baseline_simulation"]["changed_from_source_count"] == 0
        and integration["api_smoke"]["status"] == "PASS"
        and evaluation["ground_truth_leakage_audit"]["pass"]
        and evaluation["metrics"]["affected_task_retrieval"]["l2_exclusion_audit"]["pass"]
    )
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
